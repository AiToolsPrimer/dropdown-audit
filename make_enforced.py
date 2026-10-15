#!/usr/bin/env python3
"""Two drop-downs made with openpyxl: the documented example, and one that enforces the list."""
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Tickets'
ws['A1'], ws['B1'] = 'Default (accepts anything)', 'Enforced'

# As in the openpyxl docs: a message is set, but showErrorMessage stays at its default, False.
loose = DataValidation(type='list', formula1='"Open,Closed,Waiting"', allow_blank=True)
loose.error = 'Your entry is not in the list'

# The same list, with the alert switched on and the Stop style written explicitly.
strict = DataValidation(type='list', formula1='"Open,Closed,Waiting"', allow_blank=True,
                        showErrorMessage=True, errorStyle='stop',
                        errorTitle='Not on the list', error='Pick a value from the drop-down.')

for dv, cells in ((loose, 'A2:A100'), (strict, 'B2:B100')):
    ws.add_data_validation(dv)
    dv.add(cells)

wb.save('dropdowns.xlsx')
print('saved dropdowns.xlsx')
